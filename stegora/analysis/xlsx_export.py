"""
Stegora XLSX Export Module

Exports testing matrix results to Excel (XLSX) format.
Provides formatted spreadsheet with test results and summary statistics.

Academic Context:
- Professional reporting format for academic submission
- Clear presentation of experimental results
- Easy to analyze and compare metrics

Excel Structure:
- Sheet 1: Test Results (detailed 15 cases)
- Sheet 2: Summary Statistics
- Formatting: Headers, colors, number formats

Author: Hana (247006111170)
Course: Information Security - Universitas Siliwangi
"""

from typing import List, Dict, Optional
from pathlib import Path
import io


try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class XLSXExporter:
    """
    Excel (XLSX) exporter for testing matrix results.
    
    Creates formatted Excel workbook with:
    - Detailed test results
    - Summary statistics
    - Professional formatting
    """
    
    def __init__(self):
        """Initialize XLSX exporter."""
        if not OPENPYXL_AVAILABLE:
            raise ImportError(
                "openpyxl is required for XLSX export. "
                "Install with: pip install openpyxl"
            )
        
        self.workbook = None
    
    def export_to_xlsx(
        self,
        test_cases: List[Dict],
        summary: Dict,
        output_path: str
    ) -> None:
        """
        Export test results to XLSX file.
        
        Args:
            test_cases: List of test case dictionaries
            summary: Summary statistics dictionary
            output_path: Path to output XLSX file
        
        Raises:
            ImportError: If openpyxl not available
        """
        # Create workbook
        self.workbook = openpyxl.Workbook()
        
        # Remove default sheet
        if 'Sheet' in self.workbook.sheetnames:
            del self.workbook['Sheet']
        
        # Create sheets
        self._create_results_sheet(test_cases)
        self._create_summary_sheet(summary)
        
        # Save workbook
        self.workbook.save(output_path)
    
    def _create_results_sheet(self, test_cases: List[Dict]) -> None:
        """Create detailed results sheet."""
        ws = self.workbook.create_sheet("Test Results")
        
        # Define columns
        columns = [
            ('Case ID', 'case_id', 10),
            ('Image Name', 'image_name', 20),
            ('Format', 'image_format', 10),
            ('Width', 'image_width', 10),
            ('Height', 'image_height', 10),
            ('Channels', 'image_channels', 10),
            ('Capacity (bytes)', 'usable_capacity_bytes', 15),
            ('Payload Name', 'payload_name', 20),
            ('Payload Type', 'payload_type', 12),
            ('Payload Size (bytes)', 'payload_size_bytes', 18),
            ('Utilization (%)', 'capacity_utilization', 15),
            ('Embed Success', 'embedding_success', 15),
            ('MSE', 'mse', 12),
            ('PSNR (dB)', 'psnr', 12),
            ('Extract Success', 'extraction_success', 15),
            ('Exact Match', 'exact_match', 15)
        ]
        
        # Write headers
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col_idx, (header, _, width) in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        
        # Write data
        for row_idx, test_case in enumerate(test_cases, 2):
            for col_idx, (_, key, _) in enumerate(columns, 1):
                value = test_case.get(key, '')
                
                # Format specific values
                if key == 'capacity_utilization' and value is not None:
                    value = f"{value * 100:.2f}"  # Convert to percentage
                elif key == 'mse' and value is not None:
                    value = f"{value:.4f}"
                elif key == 'psnr' and value is not None:
                    if value == float('inf'):
                        value = "∞"
                    else:
                        value = f"{value:.2f}"
                elif isinstance(value, bool):
                    value = "✓" if value else "✗"
                
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
                # Color code boolean results
                if key in ('embedding_success', 'extraction_success', 'exact_match'):
                    if test_case.get(key):
                        cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                        cell.font = Font(color="006100")
                    else:
                        cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                        cell.font = Font(color="9C0006")
        
        # Add borders
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for row in ws.iter_rows(min_row=1, max_row=len(test_cases) + 1, min_col=1, max_col=len(columns)):
            for cell in row:
                cell.border = thin_border
        
        # Freeze first row
        ws.freeze_panes = "A2"
    
    def _create_summary_sheet(self, summary: Dict) -> None:
        """Create summary statistics sheet."""
        ws = self.workbook.create_sheet("Summary")
        
        # Title
        ws['A1'] = "Testing Matrix Summary"
        ws['A1'].font = Font(bold=True, size=14)
        ws.merge_cells('A1:B1')
        
        # Summary data
        summary_data = [
            ('', ''),
            ('Total Test Cases', summary.get('total_cases', 0)),
            ('', ''),
            ('Embedding Results', ''),
            ('  Success Count', summary.get('embedding_success_count', 0)),
            ('  Success Rate', f"{summary.get('embedding_success_rate', 0) * 100:.1f}%"),
            ('', ''),
            ('Extraction Results', ''),
            ('  Success Count', summary.get('extraction_success_count', 0)),
            ('  Success Rate', f"{summary.get('extraction_success_rate', 0) * 100:.1f}%"),
            ('', ''),
            ('Exact Match', ''),
            ('  Match Count', summary.get('exact_match_count', 0)),
            ('  Match Rate', f"{summary.get('exact_match_rate', 0) * 100:.1f}%"),
            ('', ''),
            ('Quality Metrics', ''),
            ('  Average MSE', f"{summary.get('average_mse', 0):.4f}" if summary.get('average_mse') else 'N/A'),
            ('  Average PSNR (dB)', f"{summary.get('average_psnr', 0):.2f}" if summary.get('average_psnr') else 'N/A'),
            ('', ''),
            ('Capacity', ''),
            ('  Average Utilization', f"{summary.get('average_utilization', 0) * 100:.2f}%"),
        ]
        
        # Write summary
        for row_idx, (label, value) in enumerate(summary_data, 3):
            ws.cell(row=row_idx, column=1, value=label)
            ws.cell(row=row_idx, column=2, value=value)
            
            # Format headers
            if label and not label.startswith('  '):
                ws.cell(row=row_idx, column=1).font = Font(bold=True)
        
        # Set column widths
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 20
        
        # Add interpretation
        ws.cell(row=len(summary_data) + 5, column=1, value="Interpretation:")
        ws.cell(row=len(summary_data) + 5, column=1).font = Font(bold=True)
        
        interpretation = [
            "• Success Rate 100% = All operations successful",
            "• Exact Match 100% = Data recovered perfectly",
            "• PSNR ≥ 30 dB = Good visual quality",
            "• Higher Utilization = More data embedded"
        ]
        
        for idx, text in enumerate(interpretation, 1):
            ws.cell(row=len(summary_data) + 5 + idx, column=1, value=text)
            ws.merge_cells(f'A{len(summary_data) + 5 + idx}:B{len(summary_data) + 5 + idx}')


def export_matrix_to_xlsx(
    test_cases: List[Dict],
    summary: Dict,
    output_path: str
) -> None:
    """
    Convenience function to export testing matrix to XLSX.
    
    Args:
        test_cases: List of test case dictionaries
        summary: Summary statistics dictionary
        output_path: Path to output XLSX file
    
    Raises:
        ImportError: If openpyxl not available
    """
    exporter = XLSXExporter()
    exporter.export_to_xlsx(test_cases, summary, output_path)


# Export public API
__all__ = [
    'XLSXExporter',
    'export_matrix_to_xlsx',
    'OPENPYXL_AVAILABLE'
]
